"""Add a "Literature Survey" slide to the VoiceRiders deck.

Approach: duplicate the existing card-grid slide (slide 7, "Built with") so the
new slide inherits the deck's own visual template, then rewrite its text. It is
appended as slide 9, so every existing slide keeps its number and page label.

Every reference below is a real, checkable work.

Usage: python tools/add_literature_slide.py <file.pptx>
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from pptx import Presentation

path = Path(sys.argv[1])
backup = path.with_name(f"{path.stem}.BACKUP{path.suffix}")

TEMPLATE_SLIDE = 7  # 1-based: the 7-card grid ("Built with")
NEW_PAGE_NUMBER = '9'

# Shape name -> new text. Names come from the template slide, so card titles and
# their bullets stay in the same positions.
TEXT: dict[str, str] = {
    # headings
    'Text 0': 'LITERATURE SURVEY',
    'Text 1': 'Prior work and how VoiceRiders differs',

    # card 1
    'Text 4': 'Speech Recognition',
    'Text 7': 'Whisper — Radford et al., 2022',
    'Text 9': 'whisper.cpp — Gerganov (OSS)',

    # card 2
    'Text 12': 'Voice Activity Detection',
    'Text 15': 'WebRTC VAD — Google (OSS)',
    'Text 17': 'Silero VAD — Silero Team, 2021',
    'Text 19': 'Chosen here: energy-based VAD',

    # card 3
    'Text 22': 'Noise Robustness',
    'Text 25': 'Spectral subtraction — Boll, 1979',

    # card 4
    'Text 28': 'Text Normalization',
    'Text 31': 'Manning et al., 2008 — IR',
    'Text 33': 'Jurafsky & Martin (SLP, 3rd ed.)',

    # card 5
    'Text 36': 'Intent Classification',
    'Text 39': 'TF-IDF — Salton & Buckley, 1988',

    # card 6
    'Text 42': 'Fuzzy Matching',
    'Text 45': 'Levenshtein 1966 · Winkler 1990',

    # card 7
    'Text 48': 'Authorization & APIs',
    'Text 51': 'RFC 7636 — PKCE, 2015',
    'Text 53': 'Spotify Web API (2025)',

    # footer page number (the brand label "Text 54" is left as-is)
    'Text 55': NEW_PAGE_NUMBER,
}

NOTES = """\
LITERATURE SURVEY - what each work gives us, and the gap this project fills

Speech recognition   Whisper (Radford et al., 2022) shows large-scale weak
                     supervision gives robust ASR; whisper.cpp makes it run on-device.
                     Gap: neither knows anything about music commands.
Voice activity       WebRTC VAD and Silero VAD detect speech cheaply, so silence is
                     never sent to the recogniser. We use a lightweight energy VAD
                     because it costs almost nothing per audio frame.
Noise robustness     Classical spectral subtraction (Boll, 1979) cleans the signal
                     before recognition. Gap: it needs an estimate of the noise
                     spectrum, which is unreliable on a moving motorcycle.
                     VoiceRiders instead accepts noisy transcripts and recovers the
                     command downstream (normalization + fuzzy matching + context +
                     confidence scoring).
Text normalization   Manning et al. (2008) and Jurafsky & Martin define the standard
                     pipeline: tokenization, stop-word handling, normalization.
Intent classification TF-IDF (Salton & Buckley, 1988) with logistic regression is a
                     well-established, lightweight text classifier - small enough to
                     ship on a phone.
Fuzzy matching       Levenshtein (1966) and Winkler (1990) give edit-distance and
                     prefix-weighted similarity; used here to survive ASR errors.
Authorization        RFC 7636 (PKCE, 2015) lets a public mobile client authenticate
                     without a client secret; Spotify's Web API then drives playback.

THE GAP: existing work treats speech recognition, intent classification and music
playback as separate problems, and none of it targets the rider case - no wake word,
wind/engine noise, and hands-free control. VoiceRiders joins them into one pipeline
with a confidence gate that refuses to guess when it is unsure.
"""

prs = Presentation(str(path))
slides = list(prs.slides)

if len(slides) >= 9 and slides[8].shapes[0].text_frame.text.strip().upper() == 'LITERATURE SURVEY':
    print('a Literature Survey slide already exists at position 9 - nothing to do')
    raise SystemExit(0)

source = slides[TEMPLATE_SLIDE - 1]

# 1. New slide on the same layout, with the layout's placeholders removed.
dest = prs.slides.add_slide(source.slide_layout)
for shape in list(dest.shapes):
    shape._element.getparent().remove(shape._element)

# 2. Copy every shape from the template (no media/charts exist in this deck,
#    so there are no relationships to remap).
for shape in source.shapes:
    dest.shapes._spTree.append(copy.deepcopy(shape._element))

# 3. Rewrite the text, preserving each shape's formatting.
by_name = {shape.name: shape for shape in dest.shapes}
replaced = 0
for name, value in TEXT.items():
    shape = by_name.get(name)
    if shape is None or not shape.has_text_frame:
        print(f'!! template shape {name!r} not found')
        continue

    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    if not paragraph.runs:
        print(f'!! shape {name!r} has no run to edit')
        continue

    paragraph.runs[0].text = value
    for extra_run in paragraph.runs[1:]:
        extra_run.text = ''
    for extra_paragraph in frame.paragraphs[1:]:
        for extra_run in extra_paragraph.runs:
            extra_run.text = ''
    replaced += 1

dest.notes_slide.notes_text_frame.text = NOTES

if not backup.exists():
    import shutil

    shutil.copy2(path, backup)
    print(f'backup: {backup}')

prs.save(str(path))
print(f'saved:  {path}')
print(f'slides: {len(prs.slides)} (new slide appended as {len(prs.slides)})')
print(f'text shapes rewritten: {replaced}')

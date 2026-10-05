"""Attach accurate speaker notes to the VoiceRiders deck.

Speaker notes do not affect slide layout or numbering, so the evaluation results
and technical details can be added without touching a single shape.

Usage: python tools/add_notes.py <file.pptx>
NOTE: does not create a backup - the original is already preserved as
      <name>.BACKUP.pptx by tools/verify_deck.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

path = Path(sys.argv[1])

NOTES: dict[int, str] = {
    5: (
        "WHY THE CONFIDENCE ENGINE EXISTS (architecture)\n"
        "'increase volume' and 'decrease volume' are ~90% identical as strings, so fuzzy\n"
        "similarity alone must never be allowed to choose the intent. The engine combines\n"
        "intent confidence (45%), fuzzy similarity (25%), ASR confidence (20%) and context\n"
        "(10%), then applies an explicit penalty when the top two intents are a known\n"
        "confusable pair and the margin is small. Below the execute threshold the app asks\n"
        "the rider to repeat instead of guessing.\n"
        "\n"
        "An entity-driven refinement upgrades a generic PLAY to PLAY_SONG whenever a song\n"
        "number was heard, and context resolves 'play' to RESUME while playback is paused.\n"
        "\n"
        "The linguistic lexicon (stop-words, synonyms, intent vocabulary) lives in shared\n"
        "JSON files consumed by BOTH the TypeScript app and the Python training code, so\n"
        "training and on-device inference cannot drift apart."
    ),
    7: (
        "THE STACK AS ACTUALLY BUILT\n"
        "- Mobile: React Native + Expo (SDK 57), TypeScript, Expo Router. The installable\n"
        "  Android APK is produced with EAS Build.\n"
        "- NLP: Python + scikit-learn are used to TRAIN the model; the fitted TF-IDF +\n"
        "  Logistic Regression weights are exported to JSON and run on-device in\n"
        "  TypeScript (a handful of dot products, no Python at runtime).\n"
        "- Speech: on-device whisper.cpp (via whisper.rn) with a local energy VAD, so only\n"
        "  speech reaches the recogniser. Offline by default, no API key, no cloud.\n"
        "- Spotify: Web API with OAuth 2.0 Authorization Code + PKCE. No client secret is\n"
        "  used anywhere in the project.\n"
        "- Backend: a single stateless HTTPS relay (serverless, Vercel) exists only because\n"
        "  Spotify now rejects custom-scheme redirect URIs for mobile apps.\n"
        "- Storage: OS secure storage (Android Keystore / iOS Keychain) for tokens. There is\n"
        "  no database, so no SQLite."
    ),
    8: (
        "PIPELINE\n"
        "Microphone + VAD -> Whisper (on-device) -> text normalization -> intent\n"
        "classification -> entity extraction (song number) -> fuzzy matching -> context\n"
        "-> confidence gate -> command dispatcher -> Spotify Web API.\n"
        "VAD means silence is never sent to the recogniser, which saves CPU and battery.\n"
        "\n"
        "SECOND BRANCH - ADAPTIVE AUDIO MODE\n"
        "The same microphone stream is measured for ambient noise: RMS -> moving average ->\n"
        "noise zone -> hysteresis/cooldown -> small volume deltas applied ON TOP of the\n"
        "rider's current volume. A voice volume command becomes the new baseline, so the\n"
        "adaptive system never undoes it."
    ),
}

prs = Presentation(str(path))
slides = list(prs.slides)

for slide_no, text in NOTES.items():
    if slide_no > len(slides):
        print(f"!! slide {slide_no} does not exist")
        return_code = 1
        continue
    frame = slides[slide_no - 1].notes_slide.notes_text_frame
    frame.text = text
    print(f"notes added to slide {slide_no} ({len(text)} chars)")

prs.save(str(path))
print(f"saved: {path}")

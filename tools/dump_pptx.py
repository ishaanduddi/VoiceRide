"""Dump every slide's text from a .pptx so it can be reviewed.

Usage: python tools/dump_pptx.py <file.pptx> [out.txt]
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

path = Path(sys.argv[1])
out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else None
lines: list[str] = []


def emit(text: str = "") -> None:
    lines.append(text)


prs = Presentation(str(path))
emit(f"FILE: {path}")
emit(f"slide size: {prs.slide_width} x {prs.slide_height} EMU "
     f"({prs.slide_width / 914400:.2f}in x {prs.slide_height / 914400:.2f}in)")
emit(f"slides: {len(prs.slides)}")

for index, slide in enumerate(prs.slides, start=1):
    emit("")
    emit(f"================ SLIDE {index} ================")
    try:
        emit(f"layout: {slide.slide_layout.name}")
    except Exception:  # noqa: BLE001
        pass

    for shape in slide.shapes:
        label = f"shape={shape.shape_type} name='{shape.name}'"
        if shape.has_text_frame and shape.text_frame.text.strip():
            emit(f"--- {label}")
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs) or para.text
                if text.strip():
                    emit(f"    [lvl{para.level}] {text}")
        elif shape.has_table:
            emit(f"--- {label} (TABLE)")
            for row in shape.table.rows:
                emit("    | " + " | ".join(cell.text.replace("\n", " / ") for cell in row.cells) + " |")
        elif shape.shape_type == 6:  # group
            emit(f"--- {label} (GROUP)")
            for sub in shape.shapes:
                if sub.has_text_frame and sub.text_frame.text.strip():
                    for para in sub.text_frame.paragraphs:
                        if para.text.strip():
                            emit(f"      [lvl{para.level}] {para.text}")
        else:
            emit(f"--- {label} (no text)")

    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            emit(f"NOTES: {notes}")

text = "\n".join(lines)
if out_path:
    out_path.write_text(text, encoding="utf-8")
    print(f"wrote {out_path} ({len(text)} chars)")
else:
    print(text)

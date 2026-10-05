"""Report geometry + font size for text shapes, to spot overflow risks.

Usage: python tools/inspect_deck.py <file.pptx> <slide[,slide...]>
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

path = Path(sys.argv[1])
wanted = {int(x) for x in sys.argv[2].split(",")}

prs = Presentation(str(path))
EMU_PER_IN = 914400

for index, slide in enumerate(prs.slides, start=1):
    if index not in wanted:
        continue
    print(f"\n================ SLIDE {index} ================")
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue

        sizes = set()
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size is not None:
                    sizes.add(round(run.font.size.pt, 1))

        width_in = (shape.width or 0) / EMU_PER_IN
        height_in = (shape.height or 0) / EMU_PER_IN
        tf = shape.text_frame
        first = text.splitlines()[0]
        print(
            f"{shape.name:>9} | {width_in:5.2f} x {height_in:5.2f} in | "
            f"pt={sorted(sizes) or 'inherit'} | wrap={tf.word_wrap} | "
            f"autosize={tf.auto_size} | chars={len(text):3d} | lines={len(text.splitlines())} | {first[:52]!r}"
        )

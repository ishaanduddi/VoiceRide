"""Generate the 4-speaker script for the NLP deck (1 minute each).

Companion to tools/make_nlp_deck.py. Speaker names are placeholders
("Name 1" … "Name 4") so they can be filled in by hand.

Usage: python tools/make_nlp_scripts.py <out.docx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(sys.argv[1])
OUT.parent.mkdir(parents=True, exist_ok=True)

DARK = RGBColor(0x0B, 0x12, 0x20)
MUTED = RGBColor(0x5B, 0x67, 0x7A)
GREEN = RGBColor(0x14, 0x8A, 0x3E)

SPEAKERS = [
    {
        'name': 'Name 1',
        'slides': 'Slides 1–3',
        'focus': 'Introduction · why NLP is the hard part · pipeline overview',
        'script': (
            'Good morning. We built VoiceRiders, a hands-free Spotify controller for motorcycle riders. '
            'The interesting part is not the app — it is the language problem. The rider gets one attempt, '
            'with no wake word and no confirmation question, through a helmet microphone at sixty '
            'kilometres an hour. Words arrive damaged: "volum", "beleiver", "nex song". And the two '
            'commands that matter most — increase and decrease volume — are almost identical as strings, '
            'while doing the opposite is the worst possible failure.\n\n'
            'So our pipeline has six NLP stages: voice-activity detection with noise gates, Whisper '
            'speech-to-text on the phone, normalisation, intent classification, entity extraction, and a '
            'confidence engine that decides whether to act or stay silent. Everything after the ASR step '
            'is plain TypeScript, so it behaves identically in the app and in our evaluation harness.'
        ),
        'notes': (
            'Say the three constraints clearly — one shot, damaged words, opposite commands. Point at the '
            'pipeline boxes when you reach them; do not read the captions.'
        ),
    },
    {
        'name': 'Name 2',
        'slides': 'Slides 4–5',
        'focus': 'On-device ASR · normalisation · intent classification',
        'script': (
            'Speech recognition runs fully on the device using whisper.cpp with the base dot E N model, '
            'about a hundred and forty megabytes, downloaded once. No audio leaves the phone. We bias the '
            'decoder with a prompt listing our command vocabulary, which makes "next song" far more likely '
            'than "text song".\n\n'
            'The transcript then goes through normalisation: case and punctuation are stripped, phrases '
            'are rewritten — "make it louder" becomes "increase volume" — numbers become digits and '
            'stop-words are removed. Intent classification fuses two signals: a TF-IDF model over '
            'unigrams and bigrams with ten logistic-regression classes, contributing sixty-five percent, '
            'and a rule-and-fuzzy baseline contributing thirty-five. The model is trained in Python on a '
            'hundred and seventy-seven utterances and exported as JSON, so inference on the phone is a '
            'handful of dot products. On fifty-four held-out utterances the hybrid reaches zero point '
            'eight five two accuracy.'
        ),
        'notes': (
            'Pause after "no audio leaves the phone" — it is the privacy point. On the metrics strip, say '
            'only the accuracy; the table is there for questions.'
        ),
    },
    {
        'name': 'Name 3',
        'slides': 'Slides 6–7',
        'focus': 'Entity extraction · confidence fusion · anti-confusion guards',
        'script': (
            'For song requests we extract entities, not just intents. "Play believer" gives the title '
            'believer; "play believer by imagine dragons" splits on the last "by" into title and artist. '
            'Filler words are dropped, and generic phrases like "play music" are rejected because there '
            'is no song to search for. We removed number-based navigation entirely — a rider should name '
            'the song, not its position.\n\n'
            'Then comes the part we are proudest of: deciding whether to act. Confidence is a weighted '
            'blend — forty-five percent intent probability, twenty-five percent fuzzy similarity, twenty '
            'percent ASR quality, ten percent context. If the top two intents are within zero point one '
            'and are a known confusable pair, we apply a penalty, so a coin flip never executes. An '
            'explicit "up" or "down" word overrides the classifier entirely. And below the threshold the '
            'app says nothing at all — there is no "did you mean" question.'
        ),
        'notes': (
            'This is the core slide. Slow down on the four weights and on the direction guard. The worked '
            'example on the right is the one number worth reading aloud.'
        ),
    },
    {
        'name': 'Name 4',
        'slides': 'Slides 8–10',
        'focus': 'Context and song selection · noise gating · results and limits',
        'script': (
            'Context resolves ambiguity: "play" while paused means resume, "play believer" means play '
            'that song. The entity is then ranked against Spotify results — exact title, then artist '
            'match, then popularity — and playback starts inside the playlist so "next song" still works.\n\n'
            'Most apparent misunderstandings are not classification errors: they are noise that should '
            'never be classified. We gate on peak signal-to-noise, on level modulation, on a twelve-word '
            'cap and a four-and-a-half-second ceiling, and we ignore our own spoken replies.\n\n'
            'The honest result is zero point eight five two held-out accuracy, with increase versus '
            'decrease still the dominant confusion — which is exactly why the direction guard exists. The '
            'weakest class is UNKNOWN, so noise rejection is where we keep working. Thank you — we are '
            'happy to take questions.'
        ),
        'notes': (
            'End on the limitation, not the accuracy. If time is short, drop the middle paragraph and keep '
            'the numbers.'
        ),
    },
]

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('VoiceRiders — NLP Presentation Script')
run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = DARK

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = subtitle.add_run('Four speakers · one minute each · slides 1–11')
sub.font.size = Pt(12)
sub.font.color.rgb = MUTED

names = doc.add_paragraph()
names.alignment = WD_ALIGN_PARAGRAPH.CENTER
names_run = names.add_run('Name 1   ·   Name 2   ·   Name 3   ·   Name 4')
names_run.font.size = Pt(13)
names_run.font.bold = True
names_run.font.color.rgb = GREEN

doc.add_paragraph()

table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
head = table.rows[0].cells
for cell, label in zip(head, ['Speaker', 'Time', 'Slides', 'Focus']):
    cell.text = label

for item in SPEAKERS:
    row = table.add_row().cells
    row[0].text = item['name']
    row[1].text = '1 min'
    row[2].text = item['slides']
    row[3].text = item['focus']

doc.add_page_break()

for index, item in enumerate(SPEAKERS):
    heading = doc.add_paragraph()
    h = heading.add_run(f"{item['name']} · {item['slides']} · 1 minute")
    h.font.size = Pt(16)
    h.font.bold = True
    h.font.color.rgb = DARK

    focus = doc.add_paragraph()
    f = focus.add_run(item['focus'])
    f.font.size = Pt(10.5)
    f.font.color.rgb = MUTED
    f.italic = True

    body = doc.add_paragraph()
    body.paragraph_format.space_before = Pt(10)
    body.paragraph_format.line_spacing = 1.35
    body_run = body.add_run(item['script'])
    body_run.font.size = Pt(11.5)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    label = note.add_run('Delivery note: ')
    label.font.bold = True
    label.font.size = Pt(9.5)
    label.font.color.rgb = GREEN
    note_run = note.add_run(item['notes'])
    note_run.font.size = Pt(9.5)
    note_run.font.color.rgb = MUTED

    words = len(item['script'].split())
    timing = doc.add_paragraph()
    t = timing.add_run(f'Word count: {words} · about {round(words / 145 * 60)} s at a calm pace')
    t.font.size = Pt(9)
    t.font.color.rgb = MUTED

    if index < len(SPEAKERS) - 1:
        doc.add_page_break()

doc.add_page_break()
tail = doc.add_paragraph()
tr = tail.add_run('Shared notes')
tr.font.size = Pt(15)
tr.font.bold = True

for line in [
    'Total run time is four minutes: title and problem (1), ASR and classification (2), '
    'entities and confidence (3), context, noise and results (4).',
    'Do not read slide text verbatim; the slides carry the numbers, the script carries the reasoning.',
    'Demo, if allowed: say “play believer” (plays without emptying the playlist), then “next song”, '
    'then “increase volume” (phone media volume, not Spotify volume).',
    'Likely question — “why not use the cloud?” Answer: no signal on a bike, and the audio never '
    'leaves the device.',
    'Likely question — “why does it sometimes ignore you?” That is deliberate: a missed command costs '
    'a repetition, a false command changes the music at the worst moment.',
]:
    paragraph = doc.add_paragraph(line, style='List Bullet')
    paragraph.paragraph_format.line_spacing = 1.25

doc.save(OUT)
print(f'saved {OUT} ({len(SPEAKERS)} speakers)')

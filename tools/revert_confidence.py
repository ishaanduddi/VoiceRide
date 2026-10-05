"""Remove the illustrative confidence value from slide 5's example line.

The tutor asked for architecture only, so the deck should carry no
metric-looking values. The confidence ENGINE is still shown as a pipeline
component ("Context + Confidence" in the NLP Tasks list); only the number goes.

Usage: python tools/revert_confidence.py <file.pptx>
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation

path = Path(sys.argv[1])
backup = path.with_name(f"{path.stem}.BACKUP{path.suffix}")

SLIDE = 5
OLD = "Intent  =  PLAY_SONG      Entity  =  7      Confidence  =  0.94"
NEW = "Intent  =  PLAY_SONG          Entity  =  7"

prs = Presentation(str(path))
slides = list(prs.slides)
slide = slides[SLIDE - 1]

updated = 0
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    for paragraph in shape.text_frame.paragraphs:
        full = "".join(run.text for run in paragraph.runs)
        if full.strip() != OLD:
            continue

        runs = paragraph.runs
        runs[0].text = NEW
        for run in runs[1:]:
            run.text = ""
        updated += 1

if updated == 0:
    print(f"NOT FOUND on slide {SLIDE}: {OLD!r}")
    print("Nothing changed (the line may already be reverted).")
    raise SystemExit(1)

# Preserve the pristine original only if no backup exists yet.
if not backup.exists():
    shutil.copy2(path, backup)
    print(f"backup: {backup}")

prs.save(str(path))
print(f"saved:  {path}")
print(f"slide {SLIDE}: {OLD!r}\n          -> {NEW!r}")
